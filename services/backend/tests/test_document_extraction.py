from datetime import datetime, timezone
import threading
from zipfile import ZipFile

from docx import Document as WordDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest
from sqlalchemy import select, update

from app.models.canonical import Document, DocumentRevision, Job
from app.services import document_extraction as extraction
from app.worker.worker import BackgroundWorker


def make_pdf(path, pages):
    """Write valid PDF objects, including pages without a text layer."""
    writer = PdfWriter()
    font = writer._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    }))
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        if text:
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
            })
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def test_unicode_markdown_preserves_end_beyond_old_4000_character_limit(tmp_path):
    path = tmp_path / "profile.md"
    original = "# Perfil de prueba\n\n" + "Información personal de ejemplo.\n" * 400 + "Última sección: término final."
    path.write_text(original, encoding="utf-8-sig")
    result = extraction.extract_document(path, path.name, "text/markdown")
    assert result.status == "ready"
    assert result.text == original
    assert len(result.text) > 4000


@pytest.mark.parametrize("content, expected", [(b"bad\x00binary", "unsupported"), (b"caf\xe9", "needs_review"), (b" \n", "needs_review")])
def test_non_text_or_invalid_encoding_never_becomes_ready(tmp_path, content, expected):
    path = tmp_path / "notes.txt"
    path.write_bytes(content)
    result = extraction.extract_document(path, path.name, "text/plain")
    assert result.status == expected
    assert result.text is None
    assert result.detail


def test_unsupported_format_is_not_decoded_as_text(tmp_path):
    path = tmp_path / "archive.zip"
    path.write_bytes(b"PK\x03\x04This is not extracted document text")
    result = extraction.extract_document(path, path.name, "application/zip")
    assert result.status == "unsupported"
    assert result.text is None


def test_limit_rejects_without_silently_truncating(tmp_path, monkeypatch):
    path = tmp_path / "notes.txt"
    path.write_text("x" * 101)
    monkeypatch.setattr(extraction, "MAX_EXTRACTED_CHARACTERS", 100)
    result = extraction.extract_document(path, path.name, "text/plain")
    assert result.status == "needs_review"
    assert result.text is None
    assert "100 caracteres" in result.detail
    assert "No se guardó texto recortado" in result.detail


def test_real_pdf_extracts_all_pages_and_preserves_page_references(tmp_path):
    path = tmp_path / "record.pdf"
    make_pdf(path, ["First page record", "Second page final detail"])
    result = extraction.extract_document(path, path.name, "application/pdf")
    assert result.status == "ready", result.detail
    assert result.text == "[Página 1]\nFirst page record\n\n[Página 2]\nSecond page final detail"
    assert result.page_count == 2


def test_corrupt_pdf_is_failed_instead_of_fake_binary_text(tmp_path):
    path = tmp_path / "invalid.pdf"
    path.write_bytes(b"%PDF-1.4 Extracted Text: not really a valid PDF")
    result = extraction.extract_document(path, path.name, "application/pdf")
    assert result.status == "failed"
    assert result.text is None


def test_encrypted_pdf_requires_review(tmp_path):
    path = tmp_path / "private.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("example-password")
    writer.write(path)
    result = extraction.extract_document(path, path.name, "application/pdf")
    assert result.status == "needs_review"
    assert result.text is None
    assert "contraseña" in result.detail


def test_scanned_pages_without_ocr_keep_partial_text_with_honest_status(tmp_path, monkeypatch):
    path = tmp_path / "mixed.pdf"
    make_pdf(path, ["Digital first page", None])
    monkeypatch.setattr(extraction.shutil, "which", lambda _: None)
    result = extraction.extract_document(path, path.name, "application/pdf")
    assert result.status == "needs_review"
    assert result.text == "[Página 1]\nDigital first page"
    assert "Página 2" in result.detail and "tesseract" in result.detail
    assert "binario indexado" not in result.text


def test_pdf_ocr_fallback_preserves_page_number(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    make_pdf(path, [None])
    calls = []
    monkeypatch.setattr(extraction, "_ocr_language", lambda _: "spa+eng")

    def fake_ocr(file_path, page, language, deadline):
        calls.append((file_path, page, language))
        return "Constancia escaneada de ejemplo"

    monkeypatch.setattr(extraction, "_ocr_pdf_page", fake_ocr)
    result = extraction.extract_document(path, path.name, "application/pdf")
    assert result.status == "ready"
    assert result.text == "[Página 1]\nConstancia escaneada de ejemplo"
    assert result.ocr_pages == (1,)
    assert calls == [(path, 1, "spa+eng")]


def test_docx_reads_paragraphs_and_tables_in_order_with_header(tmp_path):
    path = tmp_path / "profile.docx"
    document = WordDocument()
    document.add_paragraph("Inicio de perfil")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Dato"
    table.cell(0, 1).text = "Valor de ejemplo"
    document.add_paragraph("Final del perfil")
    document.sections[0].header.paragraphs[0].text = "Encabezado de referencia"
    document.save(path)
    result = extraction.extract_document(path, path.name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert result.status == "ready", result.detail
    assert result.text.index("Inicio de perfil") < result.text.index("Dato | Valor de ejemplo") < result.text.index("Final del perfil")
    assert "Encabezado de referencia" in result.text


def test_docx_embedded_media_keeps_body_but_requires_review(tmp_path):
    path = tmp_path / "with-chart.docx"
    document = WordDocument()
    document.add_paragraph("Cuerpo de ejemplo")
    document.save(path)
    with ZipFile(path, "a") as archive:
        archive.writestr("word/charts/chart1.xml", "<chart/>")
    result = extraction.extract_document(path, path.name, "application/octet-stream")
    assert result.status == "needs_review"
    assert result.text == "Cuerpo de ejemplo"
    assert "objetos incrustados" in result.detail


async def seed_revision(db_session, path, *, deleted=False):
    doc = Document(id="test-document", title="Fixture", is_deleted=deleted)
    rev = DocumentRevision(
        id="test-revision", document_id=doc.id, storage_path=str(path),
        original_filename=path.name, mime_type="text/plain", file_size_bytes=12,
        sha256_hash="a" * 64, extraction_status="pending",
    )
    db_session.add_all([doc, rev])
    await db_session.commit()
    return rev


@pytest.mark.asyncio
async def test_worker_runs_extraction_off_event_loop_and_persists_full_text(tmp_path, db_session, monkeypatch):
    path = tmp_path / "record.txt"
    rev = await seed_revision(db_session, path)
    main_thread = threading.get_ident()

    def fake_extract(*_):
        assert threading.get_ident() != main_thread
        return extraction.ExtractionResult("ready", "dato" * 2000)

    monkeypatch.setattr("app.worker.worker.extract_document", fake_extract)
    await BackgroundWorker().handle_document_extraction(db_session, {"revision_id": rev.id})
    await db_session.refresh(rev)
    assert rev.extraction_status == "ready"
    assert len(rev.extracted_text) == 8000


@pytest.mark.asyncio
async def test_deleted_document_is_not_extracted(tmp_path, db_session, monkeypatch):
    rev = await seed_revision(db_session, tmp_path / "deleted.txt", deleted=True)

    def forbidden(*_):
        pytest.fail("A deleted document must never be extracted")

    monkeypatch.setattr("app.worker.worker.extract_document", forbidden)
    await BackgroundWorker().handle_document_extraction(db_session, {"revision_id": rev.id})
    await db_session.refresh(rev)
    assert rev.extracted_text is None
    assert rev.extraction_status == "pending"


@pytest.mark.asyncio
async def test_document_deleted_during_extraction_discards_result(tmp_path, db_session, monkeypatch):
    import app.db.session

    rev = await seed_revision(db_session, tmp_path / "deleted.txt")

    async def delete_while_extracting(*_):
        async with app.db.session.async_session_factory() as deleting_session:
            await deleting_session.execute(update(Document).values(is_deleted=True))
            await deleting_session.commit()
        return extraction.ExtractionResult("ready", "Must not be resurrected")

    monkeypatch.setattr("app.worker.worker.asyncio.to_thread", delete_while_extracting)
    result = await BackgroundWorker().handle_document_extraction(db_session, {"revision_id": rev.id})
    await db_session.refresh(rev)
    assert "descartado" in result
    assert rev.extracted_text is None


@pytest.mark.asyncio
async def test_worker_diagnostics_survive_for_unsupported_document(tmp_path, db_session):
    path = tmp_path / "archive.zip"
    path.write_bytes(b"PK")
    rev = await seed_revision(db_session, path)
    rev.mime_type = "application/zip"
    job = Job(id="extraction-job", job_type="document_extraction", payload_json={"revision_id": rev.id})
    db_session.add(job)
    await db_session.commit()
    assert await BackgroundWorker().process_single_job(db_session, job) is True
    await db_session.refresh(rev)
    assert job.status == "completed"
    assert "unsupported" in job.error_message
    assert rev.extraction_status == "unsupported"
    assert rev.extracted_text is None


@pytest.mark.asyncio
async def test_worker_atomic_claim_rejects_stale_pending_job(db_session):
    import app.db.session

    job = Job(id="claimed-job", job_type="cleanup_temp_uploads", payload_json={}, scheduled_at=datetime.now(timezone.utc).replace(tzinfo=None))
    db_session.add(job)
    await db_session.commit()
    async with app.db.session.async_session_factory() as other_session:
        stale_job = (await other_session.execute(select(Job).where(Job.id == job.id))).scalar_one()
        assert stale_job.status == "pending"
        worker = BackgroundWorker()
        assert await worker.process_single_job(db_session, job) is True
        assert await worker.process_single_job(other_session, stale_job) is None
    await db_session.refresh(job)
    assert job.attempts == 1


def test_ocr_with_gpu_gateway_success(tmp_path, monkeypatch):
    from PIL import Image
    img_path = tmp_path / "page.png"
    img = Image.new("RGB", (100, 100), color=(255, 255, 255))
    img.save(img_path)

    monkeypatch.setattr(extraction.settings, "ocr_enabled", True)
    monkeypatch.setattr(extraction.settings, "llm_api_base", "http://mock-gateway:8000/v1")
    monkeypatch.setattr(extraction.settings, "ocr_model", "glm-ocr")

    class MockResponse:
        status = 200
        def read(self):
            import json
            return json.dumps({
                "choices": [{"message": {"content": "Texto extraído por GLM-OCR"}}]
            }).encode("utf-8")
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    monkeypatch.setattr(extraction.urllib.request, "urlopen", lambda req, timeout: MockResponse())
    result = extraction._ocr_with_gpu_gateway(img_path, deadline=extraction.time.monotonic() + 10)
    assert result == "Texto extraído por GLM-OCR"


def test_ocr_with_gpu_gateway_fallback_on_error(tmp_path, monkeypatch):
    img_path = tmp_path / "page.png"
    img_path.write_bytes(b"\x89PNG\r\n\x1a\n")

    monkeypatch.setattr(extraction.settings, "ocr_enabled", True)
    monkeypatch.setattr(extraction.settings, "llm_api_base", "http://mock-gateway:8000/v1")
    monkeypatch.setattr(extraction.settings, "ocr_model", "glm-ocr")

    def mock_urlopen_fail(req, timeout):
        raise extraction.urllib.error.URLError("Connection refused")

    monkeypatch.setattr(extraction.urllib.request, "urlopen", mock_urlopen_fail)
    result = extraction._ocr_with_gpu_gateway(img_path, deadline=extraction.time.monotonic() + 10)
    assert result is None
